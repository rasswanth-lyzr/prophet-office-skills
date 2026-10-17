"""Prophet design system helpers for openpyxl workbooks.

Import from anywhere via the skill's absolute path:

    import sys; sys.path.insert(0, "/path/to/xlsx/brand")
    import prophet_xlsx as px

Spec: brand/DESIGN_SYSTEM.md. Colors are hex strings without '#'.
Fonts are written as Graphik names; Arial is the documented fallback
(Office files cannot embed a fallback chain).
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl.chart.axis import ChartLines
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (
    CharacterProperties,
    Font as DrawingFont,
    Paragraph,
    ParagraphProperties,
    RegularTextRun,
)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.indexed_list import IndexedList

ASSETS = Path(__file__).resolve().parent / "assets"
LOGO_BLACK = ASSETS / "prophet-logo-black.png"
LOGO_WHITE = ASSETS / "prophet-logo-white.png"
LOGO_RATIO = 2100 / 260  # width : height

# ---------------------------------------------------------------- palette
C = {
    "black": "000000",
    "ink": "0A0A0A",
    "white": "FFFFFF",
    "gray_light": "DADADA",
    "hairline": "E5E5E5",
    "gray_mid": "B4B4B4",
    "gray_dark": "606060",
    "blue": "334AFF",
    "green": "22FFAB",  # on black / as fill only; never text on white
    "magenta": "C214B9",
    "purple": "8037D0",
    "red": "FF0053",
    "yellow": "DEFF00",  # fill with black text only
    "light_blue": "4888FF",
    "lavender": "887BFF",
    "mint": "A4EDD6",
    "navy": "002668",
}
CHART_PALETTE = ["334AFF", "000000", "8037D0", "C214B9", "4888FF", "B4B4B4"]

# ------------------------------------------------------------------ fonts
FONT_LIGHT = "Graphik Light"  # headlines
FONT_REGULAR = "Graphik"  # body
FONT_MEDIUM = "Graphik Medium"  # labels, eyebrows, table headers
FONT_SEMIBOLD = "Graphik Semibold"  # item titles, totals
FONT_FALLBACK = "Arial"
BODY_SIZE = 10

# ----------------------------------------- financial-model color convention
INPUT_COLOR = C["blue"]  # hardcoded inputs / scenario levers (text)
FORMULA_COLOR = C["black"]  # formulas (text)
SHEET_LINK_COLOR = "008000"  # links to another sheet (standard green; 22FFAB is illegible on white)
FILE_LINK_COLOR = C["red"]  # links to another file (text)
KEY_FILL = C["yellow"]  # key assumptions / cells the user fills in (fill, black text)

NUMBER_FORMATS = {
    "currency": "$#,##0;($#,##0);-",
    "pct": "0.0%",
    "multiple": "0.0x",
    "integer": "#,##0;(#,##0);-",
}

HAIRLINE = Side(style="thin", color=C["hairline"])
RULE = Side(style="thin", color=C["black"])


def font(name: str = FONT_REGULAR, size: float = BODY_SIZE, color: str = C["ink"], **kw) -> Font:
    return Font(name=name, size=size, color=color, **kw)


def _col(c) -> int:
    return column_index_from_string(c) if isinstance(c, str) else int(c)


def _is_numeric(value) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    return isinstance(value, str) and value.startswith("=")


def _recolor(cell, color: str) -> None:
    f = cell.font
    cell.font = Font(name=f.name, size=f.size, bold=f.bold, italic=f.italic, color=color)


# --------------------------------------------------------------- workbook
def apply_defaults(wb) -> None:
    """Make Graphik 10pt the workbook default (Normal style + font 0).

    Call on a fresh workbook before styling cells; on an existing file this
    would restyle every unstyled cell — don't (its conventions win).
    """
    base = Font(name=FONT_REGULAR, size=BODY_SIZE, color=C["ink"], family=2, scheme=None)
    fonts = IndexedList([base])
    for f in list(wb._fonts)[1:]:
        fonts.add(f)
    wb._fonts = fonts
    for style in wb._named_styles:
        if style.name == "Normal":
            style.font = base
    for ws in wb.worksheets:
        ws.sheet_format.defaultRowHeight = 15


# ------------------------------------------------------------------ sheet
def title_block(ws, title: str, subtitle: str | None = None, eyebrow: str | None = None,
                start_row: int = 1, col: int | str = 1) -> int:
    """Eyebrow / title / subtitle at the top of a sheet. Returns next free row."""
    c = _col(col)
    r = start_row
    if eyebrow:
        cell = ws.cell(row=r, column=c, value=eyebrow.upper())
        cell.font = font(FONT_MEDIUM, 9, C["blue"])
        ws.row_dimensions[r].height = 16
        r += 1
    cell = ws.cell(row=r, column=c, value=title)
    cell.font = font(FONT_LIGHT, 20, C["black"])
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[r].height = 32
    r += 1
    if subtitle:
        cell = ws.cell(row=r, column=c, value=subtitle)
        cell.font = font(FONT_REGULAR, 10, C["gray_dark"])
        r += 1
    return r + 1  # one blank spacer row


def section_label(ws, cell_ref: str, text: str) -> None:
    """Small uppercase blue eyebrow above a table."""
    ws[cell_ref] = text.upper()
    ws[cell_ref].font = font(FONT_MEDIUM, 9, C["blue"])


def header_row(ws, row: int, first_col, last_col, numeric_cols=None, height: float = 22) -> None:
    """Black fill, white Graphik Medium 10pt. Numeric columns right-aligned.

    numeric_cols: iterable of column letters/indices; if None, a column counts
    as numeric when the cell below the header holds a number or a formula.
    """
    a, b = _col(first_col), _col(last_col)
    numeric = {_col(x) for x in numeric_cols} if numeric_cols is not None else None
    fill = PatternFill("solid", fgColor=C["black"])
    for c in range(a, b + 1):
        cell = ws.cell(row=row, column=c)
        is_num = (c in numeric) if numeric is not None else _is_numeric(ws.cell(row=row + 1, column=c).value)
        cell.fill = fill
        cell.font = font(FONT_MEDIUM, 10, C["white"])
        cell.alignment = Alignment(horizontal="right" if is_num else "left", vertical="center", indent=0 if is_num else 1)
        cell.border = Border()
    ws.row_dimensions[row].height = height


def body_style(ws, cell_range: str, number_format: str | None = None) -> None:
    """Graphik 10pt ink, hairline E5E5E5 bottom rule, no vertical rules, no zebra."""
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            num = _is_numeric(cell.value)
            cell.font = font()
            cell.border = Border(bottom=HAIRLINE)
            cell.alignment = Alignment(horizontal="right" if num else "left", vertical="center", indent=0 if num else 1)
            if number_format and num:
                cell.number_format = number_format
        ws.row_dimensions[row[0].row].height = 18


def total_row(ws, row: int, first_col, last_col, number_format: str | None = None) -> None:
    """Graphik Semibold, thin black top rule."""
    for c in range(_col(first_col), _col(last_col) + 1):
        cell = ws.cell(row=row, column=c)
        num = _is_numeric(cell.value)
        f = cell.font
        cell.font = font(FONT_SEMIBOLD, BODY_SIZE, f.color.rgb[-6:] if f.color is not None and isinstance(f.color.rgb, str) else C["black"])
        cell.border = Border(top=RULE)
        cell.alignment = Alignment(horizontal="right" if num else "left", vertical="center", indent=0 if num else 1)
        if number_format and num:
            cell.number_format = number_format
    ws.row_dimensions[row].height = 20


_SHEET_LINK = re.compile(r"^=\s*('[^']+'|[A-Za-z0-9_.]+)!\$?[A-Z]{1,3}\$?\d+\s*$")


def apply_model_colors(ws, cell_range: str) -> None:
    """Financial-model text colors by content: numeric hardcode → INPUT_COLOR,
    formula → FORMULA_COLOR, pure cross-sheet reference (='Sheet'!B5) →
    SHEET_LINK_COLOR, any external-file reference ('[n]') → FILE_LINK_COLOR.
    Text labels are left alone. Call after body_style/total_row."""
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                if "[" in v:
                    color = FILE_LINK_COLOR
                elif _SHEET_LINK.match(v):
                    color = SHEET_LINK_COLOR
                else:
                    color = FORMULA_COLOR
            elif _is_numeric(v):
                color = INPUT_COLOR
            else:
                continue
            _recolor(cell, color)


def mark_key(cell) -> None:
    """Key assumption / user-fill cell: DEFF00 fill, black text."""
    cell.fill = PatternFill("solid", fgColor=KEY_FILL)
    _recolor(cell, C["black"])


def note(ws, cell_ref: str, text: str) -> None:
    """Caption / source line: Graphik 8pt gray-dark."""
    ws[cell_ref] = text
    ws[cell_ref].font = font(FONT_REGULAR, 8, C["gray_dark"])


def add_wordmark(ws, anchor: str = "A1", height_px: int = 16, white: bool = False) -> None:
    """Place the PROPHET wordmark (needs Pillow). Never stretched."""
    from openpyxl.drawing.image import Image

    img = Image(str(LOGO_WHITE if white else LOGO_BLACK))
    img.height = height_px
    img.width = round(height_px * LOGO_RATIO)
    ws.add_image(img, anchor)


def finish_sheet(ws, freeze: str | None = None, widths: dict | None = None,
                 min_width: float = 8, max_width: float = 48, measure_from_row: int = 1,
                 tab_color: str = C["black"], landscape: bool = False, one_page: bool = False) -> None:
    """Hide gridlines, freeze panes, size columns, tab color, Prophet print footer.

    Auto widths measure cells from `measure_from_row` down (start below the
    title block so a long title doesn't widen column A); `widths` overrides.
    Widths include ~10% slack because Graphik runs wider than Arial.
    """
    ws.sheet_view.showGridLines = False
    if freeze:
        ws.freeze_panes = freeze
    ws.sheet_properties.tabColor = tab_color

    best: dict[int, float] = {}
    for row in ws.iter_rows(min_row=measure_from_row):
        for cell in row:
            if cell.value is None or type(cell).__name__ == "MergedCell":
                continue
            v = cell.value
            if isinstance(v, str) and v.startswith("=") or isinstance(v, float):
                n = 10
            else:
                n = len(str(v))
            best[cell.column] = max(best.get(cell.column, 0), n)
    for col, n in best.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, n * 1.1 + 3))
    for key, w in (widths or {}).items():
        ws.column_dimensions[key if isinstance(key, str) else get_column_letter(key)].width = w

    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if one_page else 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    for part, text in ((ws.oddFooter.left, "PROPHET"), (ws.oddFooter.right, "&P")):
        part.text = text
        part.font = f"{FONT_MEDIUM},Regular"
        part.size = 8
        part.color = C["gray_dark"]


# ------------------------------------------------------------------ chart
def _rich(size: float = 9, color: str = C["gray_dark"], face: str = FONT_REGULAR, bold: bool = False) -> RichText:
    cp = CharacterProperties(sz=int(size * 100), b=bold, solidFill=color,
                             latin=DrawingFont(typeface=face), cs=DrawingFont(typeface=face))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def chart_title(text: str, size: float = 12, color: str = C["black"]) -> Title:
    cp = CharacterProperties(sz=int(size * 100), b=False, solidFill=color,
                             latin=DrawingFont(typeface=FONT_LIGHT), cs=DrawingFont(typeface=FONT_LIGHT))
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def style_chart(chart, title: str | None = None, data_labels: bool = True,
                number_format: str | None = None, palette=CHART_PALETTE) -> None:
    """Flat Prophet chart: palette series, hairline gridlines, no border,
    Graphik 9pt 606060 axis text, data labels on."""
    for i, s in enumerate(chart.series):
        color = palette[i % len(palette)]
        gp = s.graphicalProperties or GraphicalProperties()
        if getattr(chart, "tagname", "") in ("lineChart", "scatterChart"):
            gp.line = LineProperties(solidFill=color, w=28575)
            s.smooth = False
        else:
            gp.solidFill = color
            gp.line = LineProperties(noFill=True)
        s.graphicalProperties = gp
        if data_labels:
            s.dLbls = DataLabelList(showVal=True, showSerName=False, showCatName=False,
                                    showLegendKey=False, showPercent=False, txPr=_rich(8, C["ink"]))
            if number_format:
                s.dLbls.numFmt = number_format

    for axis in (getattr(chart, "x_axis", None), getattr(chart, "y_axis", None)):
        if axis is None:
            continue
        axis.delete = False
        axis.txPr = _rich()
        axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=C["hairline"], w=6350))
        axis.majorTickMark = "none"
        axis.minorTickMark = "none"
        axis.majorGridlines = None
    y = getattr(chart, "y_axis", None)
    if y is not None:
        y.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=C["hairline"], w=6350)))
        y.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
        if number_format:
            y.number_format = number_format

    if getattr(chart, "tagname", "") == "barChart":
        chart.gapWidth = 60
        chart.overlap = 0 if chart.grouping == "clustered" else 100
    if chart.legend is not None:
        chart.legend.position = "b"
        chart.legend.txPr = _rich()
        if len(chart.series) == 1:
            chart.legend = None
    if title:
        chart.title = chart_title(title)
    chart.style = None
    chart.roundedCorners = False
    chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    chart.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
