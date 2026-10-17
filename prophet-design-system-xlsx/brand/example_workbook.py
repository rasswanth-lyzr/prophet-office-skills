"""Build a short Prophet-branded workbook with the brand helper.

    python brand/example_workbook.py [output.xlsx]     # default ./prophet-example.xlsx
    python scripts/recalc.py output.xlsx               # then recalculate (mandatory)

All figures are illustrative.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import BarChart, Reference  # noqa: E402

import prophet_xlsx as px  # noqa: E402

out = sys.argv[1] if len(sys.argv) > 1 else "./prophet-example.xlsx"

wb = Workbook()
px.apply_defaults(wb)
ws = wb.active
ws.title = "Summary"

px.add_wordmark(ws, "A1")
r = px.title_block(
    ws,
    "Digital experience program outlook",
    subtitle="Five-year revenue and contribution view. Illustrative figures, not client data.",
    eyebrow="Experience & Innovation",
    start_row=3,
)  # -> 7

# ---------------------------------------------------------------- inputs
px.section_label(ws, f"A{r}", "Inputs")
hdr = r + 1  # 8
inputs = [
    ("Base-year revenue ($k)", 4200, px.NUMBER_FORMATS["integer"], "Illustrative"),
    ("Annual revenue growth", 0.18, px.NUMBER_FORMATS["pct"], "Key lever"),
    ("Gross margin", 0.42, px.NUMBER_FORMATS["pct"], "Illustrative"),
    ("Program investment, year 1 ($k)", 650, px.NUMBER_FORMATS["integer"], "Illustrative"),
    ("Investment step-up per year", 0.05, px.NUMBER_FORMATS["pct"], "Illustrative"),
]
for c, v in enumerate(["Assumption", "Value", "Basis"], start=1):
    ws.cell(row=hdr, column=c, value=v)
for i, (label, value, fmt, basis) in enumerate(inputs, start=hdr + 1):
    ws.cell(row=i, column=1, value=label)
    ws.cell(row=i, column=2, value=value).number_format = fmt
    ws.cell(row=i, column=3, value=basis)
first_in, last_in = hdr + 1, hdr + len(inputs)  # 9..13
px.header_row(ws, hdr, "A", "C")
px.body_style(ws, f"A{first_in}:C{last_in}")
px.apply_model_colors(ws, f"B{first_in}:B{last_in}")
px.mark_key(ws[f"B{first_in + 1}"])  # growth = key assumption
px.note(ws, f"A{last_in + 1}",
        "Blue text = input. Black = formula. Yellow fill = key assumption. All values illustrative.")

BASE, GROWTH, MARGIN, INVEST, STEP = (f"$B${first_in + k}" for k in range(5))

# ------------------------------------------------------------ projection
top = last_in + 3  # 16
px.section_label(ws, f"A{top}", "Projection ($k), illustrative")
ph = top + 1  # 17
years = ["2026", "2027", "2028", "2029", "2030"]  # years as text
cols = "BCDEF"
ws.cell(row=ph, column=1, value="Metric")
for c, y in zip(cols, years):
    ws[f"{c}{ph}"] = y
ws[f"G{ph}"] = "Total"

rev, gp, inv, net, mgn = ph + 1, ph + 2, ph + 3, ph + 4, ph + 5
ws[f"A{rev}"], ws[f"A{gp}"], ws[f"A{inv}"] = "Revenue", "Gross profit", "Program investment"
ws[f"A{net}"], ws[f"A{mgn}"] = "Net contribution", "Contribution margin"
for k, c in enumerate(cols):
    prev = cols[k - 1] if k else None
    ws[f"{c}{rev}"] = f"={BASE}" if k == 0 else f"={prev}{rev}*(1+{GROWTH})"
    ws[f"{c}{gp}"] = f"={c}{rev}*{MARGIN}"
    ws[f"{c}{inv}"] = f"={INVEST}" if k == 0 else f"={prev}{inv}*(1+{STEP})"
    ws[f"{c}{net}"] = f"={c}{gp}-{c}{inv}"
for row in (rev, gp, inv, net):
    ws[f"G{row}"] = f"=SUM(B{row}:F{row})"
for c in cols + "G":
    ws[f"{c}{mgn}"] = f"=IFERROR({c}{net}/{c}{rev},0)"

px.header_row(ws, ph, "A", "G")
px.body_style(ws, f"A{rev}:G{inv}", number_format=px.NUMBER_FORMATS["currency"])
px.total_row(ws, net, "A", "G", number_format=px.NUMBER_FORMATS["currency"])
px.body_style(ws, f"A{mgn}:G{mgn}", number_format=px.NUMBER_FORMATS["pct"])
px.apply_model_colors(ws, f"B{rev}:G{mgn}")

# ----------------------------------------------------------------- chart
chart = BarChart()
chart.type = "col"
chart.grouping = "clustered"
chart.add_data(Reference(ws, min_col=1, max_col=6, min_row=rev), from_rows=True, titles_from_data=True)
chart.add_data(Reference(ws, min_col=1, max_col=6, min_row=net), from_rows=True, titles_from_data=True)
chart.set_categories(Reference(ws, min_col=2, max_col=6, min_row=ph))
px.style_chart(chart, title="Revenue and net contribution ($k)", number_format="#,##0")
chart.width, chart.height = 17.5, 8
ws.add_chart(chart, f"A{mgn + 3}")
px.note(ws, f"A{mgn + 20}", "Source: Illustrative model for demonstration only.")

px.finish_sheet(ws, freeze=f"A{r}", measure_from_row=r + 1, widths={"A": 34}, landscape=True, one_page=True)
wb.save(out)
print(out)
